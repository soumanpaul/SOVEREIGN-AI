"""Run presentation-v1.0.0 through the live SovereignForgeAI HTTP API."""

from __future__ import annotations

import hashlib
import json
import os
import time
from pathlib import Path
from typing import Any

import httpx
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "output"
API = os.environ.get("SOVEREIGN_API_URL", "http://localhost:8000/api/v1").rstrip("/")
EMAIL = os.environ.get(
    "PRESENTATION_DEMO_EMAIL", "presentation-v1@sovereignforge.local"
)
PASSWORD = os.environ.get("PRESENTATION_DEMO_PASSWORD")
TERMINAL = {"completed", "failed", "cancelled", "timed_out"}


def require(response: httpx.Response, *statuses: int) -> Any:
    if response.status_code not in statuses:
        raise RuntimeError(
            f"{response.request.method} {response.request.url}: "
            f"HTTP {response.status_code} {response.text[:1000]}"
        )
    if response.status_code == 204:
        return None
    return response.json()


def sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def authenticate(client: httpx.Client) -> dict[str, Any]:
    if not PASSWORD:
        raise RuntimeError("Set PRESENTATION_DEMO_PASSWORD before running acceptance.")
    signup = client.post(
        f"{API}/auth/signup",
        json={
            "organization_name": "Aegis Process Systems Pvt. Ltd.",
            "full_name": "Aegis Demo Owner",
            "email": EMAIL,
            "password": PASSWORD,
        },
    )
    if signup.status_code == 201:
        return signup.json()
    if signup.status_code != 409:
        require(signup, 201)
    return require(
        client.post(
            f"{API}/auth/signin", json={"email": EMAIL, "password": PASSWORD}
        ),
        200,
    )


def ensure_workspace(client: httpx.Client, name: str) -> dict[str, Any]:
    existing = require(client.get(f"{API}/workspaces"), 200)
    match = next((item for item in existing if item["name"] == name), None)
    if match:
        return match
    return require(client.post(f"{API}/workspaces", json={"name": name}), 201)


def ensure_file(client: httpx.Client, workspace_id: str, path: Path) -> dict[str, Any]:
    digest = sha256(path.read_bytes())
    files = require(client.get(f"{API}/workspaces/{workspace_id}/files"), 200)
    match = next((item for item in files if item["sha256"] == digest), None)
    if match:
        return match
    with path.open("rb") as handle:
        return require(
            client.post(
                f"{API}/workspaces/{workspace_id}/files",
                files={"file": (path.name, handle, "application/octet-stream")},
            ),
            201,
        )


def ensure_knowledge_base(
    client: httpx.Client, workspace_id: str, name: str, file_id: str
) -> dict[str, Any]:
    bases = require(
        client.get(f"{API}/knowledge-bases", params={"workspace_id": workspace_id}),
        200,
    )
    kb = next((item for item in bases if item["name"] == name), None)
    if kb is None:
        kb = require(
            client.post(
                f"{API}/knowledge-bases",
                json={"workspace_id": workspace_id, "name": name},
            ),
            201,
        )
    current_files = require(client.get(f"{API}/workspaces/{workspace_id}/files"), 200)
    source = next(item for item in current_files if item["id"] == file_id)
    if kb["active_index_version"] is not None and source["indexed"]:
        return kb
    job = require(
        client.post(f"{API}/knowledge-bases/{kb['id']}/ingestions", json={"file_ids": [file_id]}),
        202,
    )
    deadline = time.monotonic() + 300
    while time.monotonic() < deadline:
        job = require(client.get(f"{API}/ingestions/{job['id']}"), 200)
        if job["status"] == "completed":
            refreshed = require(
                client.get(
                    f"{API}/knowledge-bases", params={"workspace_id": workspace_id}
                ),
                200,
            )
            return next(item for item in refreshed if item["id"] == kb["id"])
        if job["status"] == "failed":
            raise RuntimeError(f"Knowledge ingestion failed: {job}")
        time.sleep(1)
    raise TimeoutError("Knowledge ingestion did not reach a terminal state.")


def verify_models(client: httpx.Client) -> list[dict[str, Any]]:
    required = {"qwen3:1.7b", "qwen2.5-coder:1.5b", "nomic-embed-text", "gemma3:4b"}
    models = require(client.get(f"{API}/models"), 200)
    found = {item["model_key"] for item in models}
    if missing := required - found:
        raise RuntimeError(f"Required registered models are missing: {sorted(missing)}")
    checked = []
    for model in models:
        if model["model_key"] in required:
            checked_model = require(client.post(f"{API}/models/{model['id']}/health-check"), 200)
            if checked_model["latest_health"]["status"] != "ready":
                raise RuntimeError(f"Model is not ready: {checked_model['model_key']}")
            checked.append(checked_model)
    return checked


def submit_and_wait(
    client: httpx.Client, key: str, payload: dict[str, Any], timeout_seconds: int = 900
) -> dict[str, Any]:
    request_digest = sha256(
        json.dumps(payload, sort_keys=True, separators=(",", ":")).encode()
    )[:12]
    key = f"{key}-{request_digest}"
    accepted = require(
        client.post(f"{API}/tasks", headers={"Idempotency-Key": key}, json=payload), 202
    )
    task = get_task_with_transient_retry(client, accepted["task_id"])
    if task["status"] in {"failed", "cancelled", "timed_out"}:
        accepted = require(client.post(f"{API}/tasks/{accepted['task_id']}/retry"), 202)
    deadline = time.monotonic() + timeout_seconds
    last_status = None
    while time.monotonic() < deadline:
        task = get_task_with_transient_retry(client, accepted["task_id"])
        if task["status"] != last_status:
            print(f"{key}: {task['status']}", flush=True)
            last_status = task["status"]
        if task["status"] in TERMINAL:
            if task["status"] != "completed":
                run = task.get("latest_run") or {}
                raise RuntimeError(
                    f"{key} failed: {run.get('error_category')} {run.get('error_message')}"
                )
            return task
        time.sleep(2)
    raise TimeoutError(f"{key} did not finish within {timeout_seconds}s")


def get_task_with_transient_retry(client: httpx.Client, task_id: str) -> dict[str, Any]:
    """Tolerate a brief fixed-ingress connection recycle while local models load."""
    for attempt in range(6):
        response = client.get(f"{API}/tasks/{task_id}")
        if response.status_code == 200:
            return response.json()
        if response.status_code not in {502, 503, 504}:
            return require(response, 200)
        if attempt < 5:
            time.sleep(1 + attempt)
    return require(response, 200)


def validate_common(task: dict[str, Any], expected_type: str) -> None:
    run = task["latest_run"]
    if task["task_type"] != expected_type:
        raise AssertionError(f"Expected route {expected_type}, got {task['task_type']}")
    if not run["route"].get("selected_model_key"):
        raise AssertionError("Automatic routing did not persist a selected model.")
    if not run["steps"]:
        raise AssertionError("Execution trace is empty.")


def validate_inspection(task: dict[str, Any]) -> None:
    validate_common(task, "rag")
    run = task["latest_run"]
    text = (run["result_text"] or "").casefold()
    for term in ("p-101", "lockout", "zero", "approval"):
        if term not in text:
            raise AssertionError(f"Inspection result is missing expected concept: {term}")
    cited = {item["display_name"] for item in run["citations"]}
    required = {"aegis-p101-inspection-report.pdf", "aegis-pump-maintenance-sop.pdf"}
    if not required.issubset(cited):
        raise AssertionError(f"Inspection citations incomplete: {sorted(cited)}")


def validate_coding(task: dict[str, Any]) -> None:
    validate_common(task, "coding")
    artifacts = {item["logical_name"] for item in task["latest_run"]["artifacts"]}
    required = {"code_patch", "verified_repository", "sandbox_report"}
    if not required.issubset(artifacts):
        raise AssertionError(f"Coding artifacts incomplete: {sorted(artifacts)}")
    steps = json.dumps(task["latest_run"]["steps"]).casefold()
    for term in ("egress_blocked", "exit_code"):
        if term not in steps:
            raise AssertionError(f"Coding trace is missing evidence: {term}")


def validate_procurement(task: dict[str, Any]) -> None:
    validate_common(task, "procurement")
    text = (task["latest_run"]["result_text"] or "").casefold()
    for term in ("aravind industrial", "45,000", "beacon controls", "crest process supply"):
        if term not in text:
            raise AssertionError(f"Procurement result is missing expected value: {term}")
    extensions = {Path(item["display_name"]).suffix for item in task["latest_run"]["artifacts"]}
    if extensions != {".docx", ".xlsx"}:
        raise AssertionError(f"Procurement artifact set is invalid: {extensions}")


def download_office_artifacts(
    client: httpx.Client, workflow: str, task: dict[str, Any]
) -> list[dict[str, Any]]:
    records = []
    for artifact in task["latest_run"]["artifacts"]:
        if Path(artifact["display_name"]).suffix not in {".docx", ".xlsx"}:
            continue
        response = client.get(f"{API}{artifact['download_url'].removeprefix('/api/v1')}")
        if response.status_code != 200:
            require(response, 200)
        digest = sha256(response.content)
        if artifact["validation_status"] != "valid" or digest != artifact["sha256"]:
            raise AssertionError(
                f"Artifact validation/checksum mismatch: {artifact['display_name']}"
            )
        destination = OUTPUT / f"{workflow}-{artifact['display_name']}"
        destination.write_bytes(response.content)
        records.append(
            {
                "artifact_id": artifact["id"],
                "path": destination.relative_to(ROOT).as_posix(),
                "sha256": digest,
                "bytes": len(response.content),
                "validation_status": artifact["validation_status"],
            }
        )
    return records


def structural_office_checks(records: list[dict[str, Any]]) -> None:
    for record in records:
        path = ROOT / record["path"]
        if path.suffix == ".docx":
            document = Document(path)
            content = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                content.extend(cell.text for row in table.rows for cell in row.cells)
            for section in document.sections:
                content.extend(paragraph.text for paragraph in section.header.paragraphs)
                content.extend(paragraph.text for paragraph in section.footer.paragraphs)
            text = "\n".join(content)
            if "Aegis Process Systems Pvt. Ltd." not in text or "™" not in text:
                raise AssertionError(f"DOCX brand/trademark missing: {path.name}")
            if not text.strip():
                raise AssertionError(f"DOCX has no readable paragraphs: {path.name}")
        elif path.suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                if workbook.sheetnames != ["Recommendation", "Quotation Lines", "Policy Controls"]:
                    raise AssertionError(f"Workbook sheets invalid: {workbook.sheetnames}")
                if not str(workbook["Quotation Lines"]["E5"].value).startswith("="):
                    raise AssertionError("Workbook line-total formula is missing.")
                if workbook["Recommendation"]["B4"].value != "Aravind Industrial":
                    raise AssertionError("Workbook recommendation is not deterministic.")
            finally:
                workbook.close()


def run_procurement_only(
    client: httpx.Client,
    user: dict[str, Any],
    models: list[dict[str, Any]],
    sovereignty: dict[str, Any],
) -> None:
    """Re-run only the procurement artifact gate during focused layout QA."""
    workspace = ensure_workspace(client, "03 Pump Procurement")
    files = [
        ensure_file(client, workspace["id"], path)
        for path in (
            ROOT / "03-procurement/mechanical-seal-quotations.csv",
            ROOT / "03-procurement/aegis-procurement-policy.pdf",
        )
    ]
    task = submit_and_wait(
        client,
        "presentation-v1-procurement-layout4",
        {
            "workspace_id": workspace["id"],
            "goal": (ROOT / "03-procurement/prompt.txt").read_text().strip(),
            "mode": "auto",
            "input_file_ids": [item["id"] for item in files],
            "requested_outputs": ["xlsx", "docx"],
        },
    )
    validate_procurement(task)
    outputs = download_office_artifacts(client, "procurement", task)
    structural_office_checks(outputs)
    result = {
        "dataset": "presentation-v1.0.0",
        "scope": "procurement",
        "organization": user["organization"]["name"],
        "status": "passed",
        "sovereignty_probe": sovereignty,
        "models": [
            {"model_key": item["model_key"], "health": item["latest_health"]["status"]}
            for item in models
        ],
        "run": task_run_record(task),
        "office_outputs": outputs,
    }
    (OUTPUT / "procurement-layout-acceptance.json").write_text(
        json.dumps(result, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(result, indent=2))


def task_run_record(task: dict[str, Any]) -> dict[str, Any]:
    return {
        "task_id": task["id"],
        "run_id": task["latest_run"]["id"],
        "task_type": task["task_type"],
        "selected_model": task["latest_run"]["route"]["selected_model_key"],
        "step_count": task["latest_run"]["step_count"],
        "artifact_count": len(task["latest_run"]["artifacts"]),
    }


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    with httpx.Client(timeout=httpx.Timeout(300, connect=10), follow_redirects=True) as client:
        user = authenticate(client)
        models = verify_models(client)
        sovereignty = require(client.post(f"{API}/security/egress-test"), 200)
        if os.environ.get("PRESENTATION_DEMO_SCENARIOS") == "procurement":
            run_procurement_only(client, user, models, sovereignty)
            return

        workspaces = {
            "inspection": ensure_workspace(client, "01 Pump Inspection"),
            "coding": ensure_workspace(client, "02 Controller Fix"),
            "procurement": ensure_workspace(client, "03 Pump Procurement"),
        }
        inspection_files = {
            path.name: ensure_file(client, workspaces["inspection"]["id"], path)
            for path in (
                ROOT / "01-inspection/aegis-p101-inspection-report.pdf",
                ROOT / "01-inspection/aegis-p101-seal-leak.jpg",
                ROOT / "01-inspection/aegis-pump-maintenance-sop.pdf",
            )
        }
        kb = ensure_knowledge_base(
            client,
            workspaces["inspection"]["id"],
            "Maintenance SOPs",
            inspection_files["aegis-pump-maintenance-sop.pdf"]["id"],
        )
        coding_file = ensure_file(
            client,
            workspaces["coding"]["id"],
            ROOT / "02-coding/p101-temperature-monitor.zip",
        )
        procurement_files = [
            ensure_file(client, workspaces["procurement"]["id"], path)
            for path in (
                ROOT / "03-procurement/mechanical-seal-quotations.csv",
                ROOT / "03-procurement/aegis-procurement-policy.pdf",
            )
        ]

        inspection = submit_and_wait(
            client,
            "presentation-v1-inspection",
            {
                "workspace_id": workspaces["inspection"]["id"],
                "goal": (ROOT / "01-inspection/prompt.txt").read_text().strip(),
                "mode": "auto",
                "input_file_ids": [
                    inspection_files["aegis-p101-inspection-report.pdf"]["id"],
                    inspection_files["aegis-p101-seal-leak.jpg"]["id"],
                ],
                "knowledge_base_ids": [kb["id"]],
                "requested_outputs": ["docx"],
            },
        )
        validate_inspection(inspection)

        coding = submit_and_wait(
            client,
            "presentation-v1-coding-focused",
            {
                "workspace_id": workspaces["coding"]["id"],
                "goal": (ROOT / "02-coding/prompt.txt").read_text().strip(),
                "mode": "auto",
                "test_command": "pytest",
                "input_file_ids": [coding_file["id"]],
                "requested_outputs": ["patch", "repository", "sandbox_report"],
            },
        )
        validate_coding(coding)

        procurement = submit_and_wait(
            client,
            "presentation-v1-procurement-layout4",
            {
                "workspace_id": workspaces["procurement"]["id"],
                "goal": (ROOT / "03-procurement/prompt.txt").read_text().strip(),
                "mode": "auto",
                "input_file_ids": [item["id"] for item in procurement_files],
                "requested_outputs": ["xlsx", "docx"],
            },
        )
        validate_procurement(procurement)

        outputs = []
        outputs.extend(download_office_artifacts(client, "inspection", inspection))
        outputs.extend(download_office_artifacts(client, "procurement", procurement))
        structural_office_checks(outputs)

        result = {
            "dataset": "presentation-v1.0.0",
            "organization": user["organization"]["name"],
            "status": "passed",
            "sovereignty_probe": sovereignty,
            "models": [
                {"model_key": item["model_key"], "health": item["latest_health"]["status"]}
                for item in models
            ],
            "runs": {
                name: task_run_record(task)
                for name, task in {
                    "inspection": inspection,
                    "coding": coding,
                    "procurement": procurement,
                }.items()
            },
            "office_outputs": outputs,
        }
        (OUTPUT / "acceptance-results.json").write_text(
            json.dumps(result, indent=2) + "\n", encoding="utf-8"
        )
        print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
