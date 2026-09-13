import hashlib
import os
import uuid
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.errors import AppError
from app.procurement.types import ProcurementComparison

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass(frozen=True, slots=True)
class PublishedProcurementArtifact:
    logical_name: str
    display_name: str
    storage_key: str
    media_type: str
    size_bytes: int
    sha256: str


def _destination(
    data_root: Path, workspace_id: uuid.UUID, artifact_id: uuid.UUID, display_name: str
) -> tuple[str, Path]:
    storage_key = f"workspaces/{workspace_id}/artifacts/{artifact_id}/{display_name}"
    destination = (data_root / storage_key).resolve()
    if data_root.resolve() not in destination.parents:
        raise AppError("ARTIFACT_PATH_INVALID", "Artifact destination is invalid.", 500)
    destination.parent.mkdir(parents=True, exist_ok=True)
    return storage_key, destination


def _published(
    logical_name: str, display_name: str, storage_key: str, media_type: str, path: Path
) -> PublishedProcurementArtifact:
    content = path.read_bytes()
    return PublishedProcurementArtifact(
        logical_name,
        display_name,
        storage_key,
        media_type,
        len(content),
        hashlib.sha256(content).hexdigest(),
    )


def _brand_sheet(sheet: Worksheet, title: str, brand: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = title
    sheet["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="235D46")
    sheet["A2"] = f"{brand} · Generated locally by SOVEREIGN AI"
    sheet["A2"].font = Font(italic=True, color="647069")


def _style_table(sheet: Worksheet, header_row: int, maximum_column: int) -> None:
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C47A1C")
        cell.alignment = Alignment(vertical="center")
    sheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(maximum_column)}{sheet.max_row}"
    )
    sheet.freeze_panes = f"A{header_row + 1}"
    for column in range(1, maximum_column + 1):
        values = [
            str(sheet.cell(row, column).value or "")
            for row in range(1, sheet.max_row + 1)
        ]
        sheet.column_dimensions[get_column_letter(column)].width = min(
            48, max(12, max(map(len, values)) + 2)
        )


def create_procurement_xlsx(
    data_root: Path,
    workspace_id: uuid.UUID,
    run_id: uuid.UUID,
    comparison: ProcurementComparison,
    organization_name: str,
) -> PublishedProcurementArtifact:
    artifact_id = uuid.uuid4()
    display_name = f"procurement-comparison-{str(run_id)[:8]}.xlsx"
    storage_key, destination = _destination(
        data_root, workspace_id, artifact_id, display_name
    )
    temporary = destination.with_name(f"{destination.stem}.tmp.xlsx")
    brand = f"{organization_name.strip() or 'Organization'}™"
    try:
        workbook = Workbook()
        summary = workbook.active
        if summary is None:
            raise AppError("ARTIFACT_VALIDATION_FAILED", "XLSX has no active sheet.", 500)
        summary.title = "Recommendation"
        _brand_sheet(summary, "Procurement Recommendation", brand)
        summary.append([])
        summary.append(["Recommended vendor", comparison.recommended_vendor or "Human review"])
        summary.append(["Decision basis", comparison.recommendation_basis])
        summary.append([])
        summary.append(
            [
                "Vendor",
                "Currency",
                "Evaluated total",
                "Lead days",
                "Warranty months",
                "Status",
                "Exceptions",
            ]
        )
        for vendor in comparison.vendors:
            summary.append(
                [
                    vendor.vendor,
                    vendor.currency,
                    vendor.total,
                    vendor.maximum_lead_days,
                    vendor.minimum_warranty_months,
                    "PASS" if vendor.compliant else "REVIEW",
                    "; ".join(vendor.issues),
                ]
            )
        _style_table(summary, 7, 7)
        for cell in summary[4]:
            cell.font = Font(bold=True)
        for row in range(8, summary.max_row + 1):
            summary.cell(row, 3).number_format = '#,##0.00'
            status = summary.cell(row, 6)
            status.fill = PatternFill(
                "solid", fgColor="D9EAD3" if status.value == "PASS" else "FCE5CD"
            )

        quotes = workbook.create_sheet("Quotation Lines")
        _brand_sheet(quotes, "Normalized Quotation Lines", brand)
        quotes.append([])
        quotes.append(
            [
                "Vendor",
                "Item",
                "Quantity",
                "Unit price",
                "Line total",
                "Currency",
                "Lead days",
                "Warranty months",
                "Declared compliant",
                "Source",
                "Page",
            ]
        )
        for index, line in enumerate(comparison.lines, start=5):
            quotes.append(
                [
                    line.vendor,
                    line.item,
                    line.quantity,
                    line.unit_price,
                    f"=C{index}*D{index}",
                    line.currency,
                    line.lead_time_days,
                    line.warranty_months,
                    line.declared_compliant,
                    f"[{line.source_id}] {line.source_name}",
                    line.page,
                ]
            )
        _style_table(quotes, 4, 11)
        for row in range(5, quotes.max_row + 1):
            quotes.cell(row, 4).number_format = "#,##0.00"
            quotes.cell(row, 5).number_format = "#,##0.00"

        policy = workbook.create_sheet("Policy Controls")
        _brand_sheet(policy, "Extracted Policy Controls", brand)
        policy.append([])
        policy.append(["Control", "Extracted value"])
        policy.append(["Maximum budget", comparison.policy.maximum_budget])
        policy.append(["Maximum lead days", comparison.policy.maximum_lead_days])
        policy.append(["Minimum warranty months", comparison.policy.minimum_warranty_months])
        policy.append(["Required currency", comparison.policy.required_currency])
        _style_table(policy, 4, 2)

        workbook.properties.title = "SOVEREIGN AI Procurement Comparison"
        workbook.properties.creator = brand
        workbook.save(temporary)
        verified = load_workbook(temporary, data_only=False, read_only=True)
        if verified.sheetnames != ["Recommendation", "Quotation Lines", "Policy Controls"]:
            raise AppError(
                "ARTIFACT_VALIDATION_FAILED", "XLSX required sheets are missing.", 500
            )
        quote_sheet = verified["Quotation Lines"]
        if quote_sheet.max_row < 5 or not str(quote_sheet["E5"].value).startswith("="):
            raise AppError(
                "ARTIFACT_VALIDATION_FAILED", "XLSX quotation formulas are missing.", 500
            )
        verified.close()
        os.replace(temporary, destination)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    return _published(
        "procurement_comparison", display_name, storage_key, XLSX_MEDIA_TYPE, destination
    )


def create_procurement_docx(
    data_root: Path,
    workspace_id: uuid.UUID,
    run_id: uuid.UUID,
    goal: str,
    comparison: ProcurementComparison,
    citations: list[dict[str, object]],
    organization_name: str,
) -> PublishedProcurementArtifact:
    artifact_id = uuid.uuid4()
    display_name = f"procurement-recommendation-{str(run_id)[:8]}.docx"
    storage_key, destination = _destination(
        data_root, workspace_id, artifact_id, display_name
    )
    temporary = destination.with_suffix(".docx.tmp")
    brand = f"{organization_name.strip() or 'Organization'}™"
    try:
        document = DocxDocument()
        document.core_properties.title = "SOVEREIGN AI Procurement Recommendation"
        document.core_properties.author = brand
        styles = document.styles
        styles["Normal"].font.name = "Aptos"
        styles["Normal"].font.size = Pt(10.5)
        for style_name in ("Title", "Heading 1", "Heading 2"):
            styles[style_name].font.name = "Aptos Display"
            styles[style_name].font.color.rgb = RGBColor(35, 93, 70)
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.7)
        section.left_margin = section.right_margin = Inches(0.8)
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run(brand)
        run.bold = True
        run.font.color.rgb = RGBColor(35, 93, 70)
        title = document.add_heading("Procurement Evaluation Recommendation", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle = document.add_paragraph(f"Run {run_id} · governed local comparison")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_heading("Requested Outcome", 1)
        document.add_paragraph(goal)
        document.add_heading("Recommendation", 1)
        recommendation = (
            f"Recommend {comparison.recommended_vendor}, subject to responsible human approval."
            if comparison.recommended_vendor
            else "No automatic award recommendation. Responsible procurement review is required."
        )
        paragraph = document.add_paragraph(recommendation)
        paragraph.runs[0].bold = True
        document.add_paragraph(comparison.recommendation_basis)
        document.add_heading("Evaluated Bids", 1)
        table = document.add_table(rows=1, cols=6)
        table.style = "Light Shading Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, value in zip(
            table.rows[0].cells,
            ("Vendor", "Total", "Lead", "Warranty", "Status", "Exceptions"),
            strict=True,
        ):
            cell.text = value
            cell.paragraphs[0].runs[0].bold = True
        for vendor in comparison.vendors:
            cells = table.add_row().cells
            values = (
                vendor.vendor,
                f"{vendor.currency} {vendor.total:,.2f}",
                str(vendor.maximum_lead_days or "Unknown"),
                str(vendor.minimum_warranty_months or "Unknown"),
                "PASS" if vendor.compliant else "REVIEW",
                "; ".join(vendor.issues) or "None extracted",
            )
            for cell, value in zip(cells, values, strict=True):
                cell.text = value
        document.add_heading("Evidence and Traceability", 1)
        for citation in citations:
            page = citation.get("page_start", 1)
            document.add_paragraph(
                f"[{citation.get('source_id', 'S?')}] "
                f"{citation.get('display_name', 'Document')}, page {page}",
                style="List Bullet",
            )
        document.add_heading("Approval Controls", 1)
        controls = [
            *comparison.warnings,
            "Confirm taxes, commercial terms, vendor eligibility, and "
            "source-document authenticity.",
            "A responsible approver must review this recommendation before purchase commitment.",
        ]
        for control in controls:
            document.add_paragraph(control, style="List Bullet")
        footer = section.footer.paragraphs[0]
        footer.text = f"Confidential · {brand} · Generated locally by SOVEREIGN AI"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.save(str(temporary))
        verified = DocxDocument(str(temporary))
        headings = {paragraph.text for paragraph in verified.paragraphs}
        required = {
            "Requested Outcome",
            "Recommendation",
            "Evaluated Bids",
            "Evidence and Traceability",
            "Approval Controls",
        }
        if not required.issubset(headings) or len(verified.tables) != 1:
            raise AppError(
                "ARTIFACT_VALIDATION_FAILED", "DOCX required sections are missing.", 500
            )
        os.replace(temporary, destination)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    return _published(
        "procurement_recommendation", display_name, storage_key, DOCX_MEDIA_TYPE, destination
    )


def create_procurement_artifacts(
    data_root: Path,
    workspace_id: uuid.UUID,
    run_id: uuid.UUID,
    goal: str,
    comparison: ProcurementComparison,
    citations: list[dict[str, object]],
    organization_name: str,
) -> list[PublishedProcurementArtifact]:
    published: list[PublishedProcurementArtifact] = []
    try:
        published.append(
            create_procurement_xlsx(
                data_root, workspace_id, run_id, comparison, organization_name
            )
        )
        published.append(
            create_procurement_docx(
                data_root,
                workspace_id,
                run_id,
                goal,
                comparison,
                citations,
                organization_name,
            )
        )
    except Exception:
        for artifact in published:
            (data_root / artifact.storage_key).unlink(missing_ok=True)
        raise
    return published
