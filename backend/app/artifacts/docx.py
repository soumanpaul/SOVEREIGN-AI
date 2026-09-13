import hashlib
import os
import re
import uuid
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from docx.document import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

from app.core.errors import AppError


@dataclass(frozen=True, slots=True)
class PublishedDocument:
    display_name: str
    storage_key: str
    size_bytes: int
    sha256: str


INLINE_MARKDOWN = re.compile(
    r"(`[^`]+`|\[[^\]]+\]\([^)]+\)|\*\*.+?\*\*|__.+?__|(?<!\*)\*[^*\n]+\*(?!\*)|(?<!_)_[^_\n]+_(?!_))"
)
TABLE_DIVIDER = re.compile(r"^\s*:?-{3,}:?\s*$")


def _clean_markdown(value: str) -> str:
    """Remove model-only reasoning wrappers while preserving the final Markdown."""
    return (
        re.sub(r"<think>[\s\S]*?</think>", "", value, flags=re.IGNORECASE)
        .replace("<think>", "")
        .replace("</think>", "")
        .strip()
    )


def _add_inline(paragraph: Paragraph, text: str) -> None:
    """Translate common inline Markdown into native Word runs."""
    position = 0
    for match in INLINE_MARKDOWN.finditer(text):
        if match.start() > position:
            paragraph.add_run(
                text[position : match.start()].replace("**", "").replace("__", "")
            )
        token = match.group(0)
        if token.startswith(("**", "__")):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Aptos Mono"
            run.font.size = Pt(9)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "EAF1ED")
            run._r.get_or_add_rPr().append(shading)
        elif token.startswith("["):
            link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            label, url = link.groups() if link else (token, "")
            run = paragraph.add_run(f"{label} ({url})" if url else label)
            run.font.color.rgb = RGBColor(34, 103, 139)
            run.underline = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:].replace("**", "").replace("__", ""))


def _table_cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_table(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines) or "|" not in lines[index]:
        return False
    cells = _table_cells(lines[index + 1])
    return len(cells) >= 2 and all(TABLE_DIVIDER.fullmatch(cell) for cell in cells)


def _add_table(document: Document, rows: list[list[str]]) -> None:
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Shading Accent 1"
    for row_index, values in enumerate(rows):
        for column_index in range(width):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            value = values[column_index] if column_index < len(values) else ""
            _add_inline(cell.paragraphs[0], value)
            if row_index == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def _add_markdown(document: Document, markdown: str) -> None:
    """Render model Markdown as native Word structure instead of literal markup."""
    lines = _clean_markdown(markdown).splitlines()
    paragraph_lines: list[str] = []
    in_code = False
    code_lines: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_lines:
            paragraph = document.add_paragraph()
            _add_inline(paragraph, " ".join(line.strip() for line in paragraph_lines))
            paragraph.paragraph_format.space_after = Pt(7)
            paragraph_lines.clear()

    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush_paragraph()
            if in_code:
                paragraph = document.add_paragraph()
                run = paragraph.add_run("\n".join(code_lines))
                run.font.name = "Aptos Mono"
                run.font.size = Pt(8.5)
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EAF1ED")
                paragraph._p.get_or_add_pPr().append(shading)
                code_lines.clear()
            in_code = not in_code
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if _is_table(lines, index):
            flush_paragraph()
            rows = [_table_cells(line)]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append(_table_cells(lines[index]))
                index += 1
            _add_table(document, rows)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            paragraph = document.add_heading(level=min(len(heading.group(1)) + 1, 4))
            _add_inline(paragraph, heading.group(2))
            index += 1
            continue
        list_item = re.match(r"^(\s*)([-+*]|\d+[.)])\s+(.+)$", line)
        if list_item:
            flush_paragraph()
            ordered = list_item.group(2)[0].isdigit()
            depth = min(len(list_item.group(1).replace("\t", "    ")) // 2, 2)
            base_style = "List Number" if ordered else "List Bullet"
            style = base_style if depth == 0 else f"{base_style} {depth + 1}"
            paragraph = document.add_paragraph(style=style)
            _add_inline(paragraph, list_item.group(3))
            index += 1
            continue
        if stripped.startswith(">"):
            flush_paragraph()
            paragraph = document.add_paragraph(style="Quote")
            _add_inline(paragraph, stripped.lstrip("> "))
            index += 1
            continue
        if re.fullmatch(r"(?:---+|___+|\*\*\*+)", stripped):
            flush_paragraph()
            index += 1
            continue
        paragraph_lines.append(line)
        index += 1
    flush_paragraph()
    if code_lines:
        paragraph = document.add_paragraph()
        run = paragraph.add_run("\n".join(code_lines))
        run.font.name = "Aptos Mono"
        run.font.size = Pt(8.5)


def create_approval_docx(
    data_root: Path,
    workspace_id: uuid.UUID,
    run_id: uuid.UUID,
    goal: str,
    result: str,
    citations: list[dict[str, object]],
    organization_name: str = "Organization",
) -> PublishedDocument:
    artifact_id = uuid.uuid4()
    display_name = f"approval-recommendation-{str(run_id)[:8]}.docx"
    storage_key = f"workspaces/{workspace_id}/artifacts/{artifact_id}/{display_name}"
    destination = (data_root / storage_key).resolve()
    root = data_root.resolve()
    if root not in destination.parents:
        raise AppError("ARTIFACT_PATH_INVALID", "Artifact destination is invalid.", 500)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix(".docx.tmp")
    brand = f"{organization_name.strip() or 'Organization'}™"
    try:
        document = DocxDocument()
        document.core_properties.title = "SOVEREIGN AI Approval Recommendation"
        document.core_properties.subject = f"AI-assisted recommendation for {brand}"
        document.core_properties.author = brand
        styles = document.styles
        styles["Normal"].font.name = "Aptos"
        styles["Normal"].font.size = Pt(10.5)
        for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
            styles[style_name].font.name = "Aptos Display"
            styles[style_name].font.color.rgb = RGBColor(35, 93, 70)
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.7)
        section.left_margin = section.right_margin = Inches(0.8)
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_run = header.add_run(brand)
        header_run.bold = True
        header_run.font.size = Pt(9)
        header_run.font.color.rgb = RGBColor(35, 93, 70)
        brand_line = document.add_paragraph()
        brand_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        brand_run = brand_line.add_run(brand)
        brand_run.bold = True
        brand_run.font.size = Pt(11)
        brand_run.font.color.rgb = RGBColor(196, 122, 28)
        title = document.add_heading("SOVEREIGN AI Approval Recommendation", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle = document.add_paragraph(f"Run {run_id} · locally generated")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.runs[0].font.color.rgb = RGBColor(100, 110, 105)
        document.add_heading("Requested Outcome", level=1)
        _add_markdown(document, goal)
        document.add_heading("Recommendation", level=1)
        _add_markdown(document, result)
        document.add_heading("Evidence and Traceability", level=1)
        if citations:
            for citation in citations:
                pages = str(citation.get("page_start", 1))
                if citation.get("page_end") != citation.get("page_start"):
                    pages += f"–{citation.get('page_end')}"
                paragraph = document.add_paragraph(style="List Bullet")
                _add_inline(
                    paragraph,
                    f"**[{citation.get('source_id', 'S?')}]** "
                    f"{citation.get('display_name', 'Document')}, page {pages}",
                )
        else:
            document.add_paragraph("No knowledge-base citations were supplied for this task.")
        document.add_heading("Approval Control", level=1)
        document.add_paragraph(
            "Human review is required before operational use. This document records an AI-assisted "
            "recommendation and does not replace the responsible approver."
        )
        footer = section.footer.paragraphs[0]
        footer.text = f"Confidential · {brand} · Generated locally by SOVEREIGN AI"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(100, 110, 105)
        document.save(str(temporary))
        verified = DocxDocument(str(temporary))
        headings = {paragraph.text for paragraph in verified.paragraphs}
        required = {
            "Requested Outcome",
            "Recommendation",
            "Evidence and Traceability",
            "Approval Control",
        }
        if not required.issubset(headings):
            raise AppError("ARTIFACT_VALIDATION_FAILED", "DOCX required sections are missing.", 500)
        if "**" in "\n".join(paragraph.text for paragraph in verified.paragraphs):
            raise AppError("ARTIFACT_VALIDATION_FAILED", "DOCX contains unrendered Markdown.", 500)
        os.replace(temporary, destination)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    content = destination.read_bytes()
    return PublishedDocument(
        display_name, storage_key, len(content), hashlib.sha256(content).hexdigest()
    )
