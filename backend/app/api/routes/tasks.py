import uuid
from pathlib import Path
from typing import Annotated, Literal

from docx import Document as DocxDocument
from fastapi import APIRouter, Header
from fastapi.responses import FileResponse
from openpyxl import load_workbook
from sqlalchemy import or_, select

from app.api.dependencies import AppSettings, CurrentUser, DatabaseSession
from app.core.errors import AppError
from app.db.models import Artifact, AuditEvent, TaskRun, TaskStep, Workspace
from app.schemas.tasks import (
    ArtifactPreviewResponse,
    ArtifactResponse,
    AuditEventResponse,
    TaskAccepted,
    TaskCreate,
    TaskResponse,
    TaskStepResponse,
)
from app.services.file_storage import resolve_storage_key
from app.tasks.runtime import notify_task_worker
from app.tasks.service import (
    cancel_task,
    create_task,
    get_owned_task,
    list_tasks,
    retry_task,
    task_response,
)

router = APIRouter(tags=["tasks"])
PREVIEW_CHARS = 12_000


def _bounded_preview(value: str) -> tuple[str, bool]:
    return value[:PREVIEW_CHARS], len(value) > PREVIEW_CHARS


def _artifact_preview(
    path: Path, media_type: str
) -> tuple[Literal["document", "spreadsheet", "text", "binary"], str, bool]:
    if media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        document = DocxDocument(str(path))
        document_lines = [
            paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
        ]
        for table in document.tables[:4]:
            for table_row in table.rows[:12]:
                document_lines.append(
                    " | ".join(cell.text.replace("\n", " ") for cell in table_row.cells)
                )
        content, truncated = _bounded_preview("\n\n".join(document_lines))
        return "document", content, truncated
    if media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet_lines: list[str] = []
            for sheet in workbook.worksheets[:6]:
                sheet_lines.append(f"## {sheet.title}")
                for cell_values in sheet.iter_rows(max_row=12, max_col=12, values_only=True):
                    values = [
                        str(value) if value is not None else "" for value in cell_values
                    ]
                    if any(values):
                        sheet_lines.append(" | ".join(values))
            content, truncated = _bounded_preview("\n".join(sheet_lines))
            return "spreadsheet", content, truncated
        finally:
            workbook.close()
    if media_type.startswith("text/") or media_type == "application/json":
        content, truncated = _bounded_preview(path.read_text(encoding="utf-8", errors="replace"))
        return "text", content, truncated
    return (
        "binary",
        "Preview is unavailable for this binary artifact. Download it for review.",
        False,
    )


@router.post("/tasks", response_model=TaskAccepted, status_code=202)
def submit_task(
    payload: TaskCreate,
    session: DatabaseSession,
    user: CurrentUser,
    settings: AppSettings,
    idempotency_key: Annotated[str | None, Header(max_length=120)] = None,
) -> TaskAccepted:
    accepted = create_task(session, payload, user, settings, idempotency_key)
    notify_task_worker()
    return accepted


@router.get("/tasks", response_model=list[TaskResponse])
def tasks(session: DatabaseSession, user: CurrentUser, limit: int = 25) -> list[TaskResponse]:
    return list_tasks(session, user, max(1, min(limit, 100)))


@router.get("/tasks/{task_id}", response_model=TaskResponse)
def task_detail(task_id: uuid.UUID, session: DatabaseSession, user: CurrentUser) -> TaskResponse:
    return task_response(session, get_owned_task(session, task_id, user))


@router.post("/tasks/{task_id}/cancel", response_model=TaskResponse, status_code=202)
def cancel(task_id: uuid.UUID, session: DatabaseSession, user: CurrentUser) -> TaskResponse:
    return cancel_task(session, get_owned_task(session, task_id, user))


@router.post("/tasks/{task_id}/retry", response_model=TaskAccepted, status_code=202)
def retry(
    task_id: uuid.UUID,
    session: DatabaseSession,
    user: CurrentUser,
    settings: AppSettings,
) -> TaskAccepted:
    accepted = retry_task(session, get_owned_task(session, task_id, user), settings)
    notify_task_worker()
    return accepted


@router.get("/runs/{run_id}/steps", response_model=list[TaskStepResponse])
def run_steps(
    run_id: uuid.UUID, session: DatabaseSession, user: CurrentUser
) -> list[TaskStep]:
    run = session.get(TaskRun, run_id)
    if run is None:
        raise AppError("RUN_NOT_FOUND", "Run not found.", 404)
    get_owned_task(session, run.task_id, user)
    return list(
        session.scalars(
            select(TaskStep).where(TaskStep.run_id == run_id).order_by(TaskStep.sequence)
        )
    )


@router.get("/runs/{run_id}/artifacts", response_model=list[ArtifactResponse])
def run_artifacts(
    run_id: uuid.UUID, session: DatabaseSession, user: CurrentUser
) -> list[ArtifactResponse]:
    run = session.get(TaskRun, run_id)
    if run is None:
        raise AppError("RUN_NOT_FOUND", "Run not found.", 404)
    get_owned_task(session, run.task_id, user)
    return [
        ArtifactResponse.model_validate(item).model_copy(
            update={"download_url": f"/api/v1/artifacts/{item.id}/content"}
        )
        for item in session.scalars(select(Artifact).where(Artifact.run_id == run_id))
    ]


@router.get("/artifacts/{artifact_id}/content")
def artifact_content(
    artifact_id: uuid.UUID,
    session: DatabaseSession,
    user: CurrentUser,
    settings: AppSettings,
) -> FileResponse:
    artifact = session.scalar(
        select(Artifact)
        .join(Workspace, Workspace.id == Artifact.workspace_id)
        .where(Artifact.id == artifact_id, Workspace.organization_id == user.organization_id)
    )
    if artifact is None:
        raise AppError("ARTIFACT_NOT_FOUND", "Artifact not found.", 404)
    path: Path = resolve_storage_key(settings.data_root, artifact.storage_key)
    if not path.is_file():
        raise AppError("ARTIFACT_FILE_MISSING", "Artifact content is unavailable.", 404)
    return FileResponse(path, media_type=artifact.media_type, filename=artifact.display_name)


@router.get("/artifacts/{artifact_id}/preview", response_model=ArtifactPreviewResponse)
def artifact_preview(
    artifact_id: uuid.UUID,
    session: DatabaseSession,
    user: CurrentUser,
    settings: AppSettings,
) -> ArtifactPreviewResponse:
    artifact = session.scalar(
        select(Artifact)
        .join(Workspace, Workspace.id == Artifact.workspace_id)
        .where(Artifact.id == artifact_id, Workspace.organization_id == user.organization_id)
    )
    if artifact is None:
        raise AppError("ARTIFACT_NOT_FOUND", "Artifact not found.", 404)
    path = resolve_storage_key(settings.data_root, artifact.storage_key)
    if not path.is_file():
        raise AppError("ARTIFACT_FILE_MISSING", "Artifact content is unavailable.", 404)
    preview_type, content, truncated = _artifact_preview(path, artifact.media_type)
    return ArtifactPreviewResponse(
        artifact_id=artifact.id,
        display_name=artifact.display_name,
        media_type=artifact.media_type,
        preview_type=preview_type,
        content=content,
        truncated=truncated,
    )


@router.get("/audit-events", response_model=list[AuditEventResponse])
def audit_events(
    session: DatabaseSession,
    user: CurrentUser,
    run_id: uuid.UUID | None = None,
    event_type: str | None = None,
    security_only: bool = False,
    limit: int = 100,
) -> list[AuditEvent]:
    statement = (
        select(AuditEvent)
        .join(Workspace, Workspace.id == AuditEvent.workspace_id)
        .where(Workspace.organization_id == user.organization_id)
        .order_by(AuditEvent.occurred_at.desc())
        .limit(max(1, min(limit, 200)))
    )
    if run_id:
        statement = statement.where(AuditEvent.run_id == run_id)
    if event_type:
        statement = statement.where(AuditEvent.event_type == event_type[:80])
    if security_only:
        statement = statement.where(
            or_(
                AuditEvent.event_type.like("%DENIED%"),
                AuditEvent.event_type.like("%FAILED%"),
                AuditEvent.event_type.like("%EGRESS%"),
                AuditEvent.event_type.like("%CANCELLED%"),
                AuditEvent.event_type.like("%TIMED_OUT%"),
            )
        )
    return list(session.scalars(statement))
