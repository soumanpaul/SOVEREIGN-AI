import uuid
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook

from app.artifacts.procurement import create_procurement_artifacts
from app.core.errors import AppError
from app.services.procurement import compare_procurement, comparison_markdown


def evidence() -> list[dict[str, object]]:
    return [
        {
            "source_id": "S1",
            "display_name": "quotations.csv",
            "page_start": 1,
            "text": (
                "vendor,item,quantity,unit_price,currency,lead_time_days,"
                "warranty_months,compliant\n"
                "Aravind Industrial,Pressure sensor,10,4500,INR,12,24,yes\n"
                "Beacon Controls,Pressure sensor,10,4200,INR,20,12,yes"
            ),
        },
        {
            "source_id": "S2",
            "display_name": "policy.md",
            "page_start": 1,
            "text": (
                "Maximum budget: INR 50,000\n"
                "Maximum lead time: 15 days\n"
                "Minimum warranty: 18 months\n"
                "Required currency: INR"
            ),
        },
    ]


def test_procurement_comparison_is_deterministic_and_policy_aware() -> None:
    comparison = compare_procurement(evidence())

    assert comparison.recommended_vendor == "Aravind Industrial"
    assert comparison.policy.maximum_budget == 50_000
    assert comparison.vendors[0].total == 45_000
    beacon = next(item for item in comparison.vendors if item.vendor == "Beacon Controls")
    assert beacon.total == 42_000
    assert beacon.compliant is False
    assert len(beacon.issues) == 2


def test_procurement_rejects_unstructured_quotes() -> None:
    with pytest.raises(AppError) as error:
        compare_procurement(
            [{"source_id": "S1", "display_name": "note.txt", "text": "Call vendor."}]
        )
    assert error.value.code == "PROCUREMENT_QUOTES_REQUIRED"


def test_governed_result_always_requires_human_review() -> None:
    result = comparison_markdown(compare_procurement(evidence()))

    assert "subject to human review" in result
    assert "No human review required" not in result


def test_procurement_artifacts_are_branded_and_structurally_valid(
    tmp_path: Path,
) -> None:
    workspace_id, run_id = uuid.uuid4(), uuid.uuid4()
    artifacts = create_procurement_artifacts(
        tmp_path,
        workspace_id,
        run_id,
        "Compare quotations.",
        compare_procurement(evidence()),
        [
            {
                "source_id": "S1",
                "display_name": "quotations.csv",
                "page_start": 1,
            }
        ],
        "Acme Engineering",
    )

    assert {item.logical_name for item in artifacts} == {
        "procurement_comparison",
        "procurement_recommendation",
    }
    xlsx = next(item for item in artifacts if item.display_name.endswith(".xlsx"))
    workbook = load_workbook(tmp_path / xlsx.storage_key, data_only=False)
    assert workbook.sheetnames == ["Recommendation", "Quotation Lines", "Policy Controls"]
    assert workbook["Recommendation"]["B4"].value == "Aravind Industrial"
    assert workbook["Quotation Lines"]["E5"].value == "=C5*D5"
    assert "Acme Engineering™" in workbook["Recommendation"]["A2"].value
    workbook.close()
    docx = next(item for item in artifacts if item.display_name.endswith(".docx"))
    document = Document(tmp_path / docx.storage_key)
    assert document.core_properties.author == "Acme Engineering™"
    assert any("Aravind Industrial" in paragraph.text for paragraph in document.paragraphs)
    assert len(document.tables) == 1
