import uuid
from pathlib import Path

import pymupdf
import pytest
from docx import Document as DocxDocument
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.artifacts.docx import create_approval_docx
from app.core.config import Settings
from app.core.errors import AppError
from app.db.base import Base
from app.db.models import (
    Document,
    KnowledgeBase,
    Organization,
    RegisteredModel,
    StoredFile,
    Task,
    TaskRun,
    User,
    Workspace,
)
from app.model_providers.base import EmbeddingResult
from app.model_providers.ollama import OllamaModelProvider
from app.routing.router import classify_task
from app.schemas.common import ModelCreate
from app.schemas.tasks import TaskCreate
from app.services.hybrid_retrieval import gather_hybrid_evidence
from app.services.model_registry import get_model, list_models, register_model, set_model_enabled
from app.tasks.service import create_task, get_owned_task
from app.tools.registry import ToolContext, ToolRegistry


@pytest.fixture
def session() -> Session:
    engine = create_engine("sqlite+pysqlite://")
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False) as value:
        yield value
    Base.metadata.drop_all(engine)
    engine.dispose()


def tenant(session: Session, name: str) -> tuple[User, Workspace]:
    organization = Organization(name=name, slug=f"{name.casefold()}-{uuid.uuid4().hex[:6]}")
    session.add(organization)
    session.flush()
    user = User(
        organization_id=organization.id,
        email=f"{uuid.uuid4().hex}@example.test",
        full_name=f"{name} Owner",
        password_hash="not-used",
    )
    workspace = Workspace(organization_id=organization.id, name=f"{name} Workspace")
    session.add_all([user, workspace])
    session.commit()
    return user, workspace


class FakeEmbeddingProvider:
    async def embed(self, texts: list[str], model_key: str) -> EmbeddingResult:
        vectors = [
            [1.0, 0.0]
            if any(word in text.casefold() for word in ("pump", "isolation"))
            else [0.0, 1.0]
            for text in texts
        ]
        return EmbeddingResult(vectors=vectors, duration_ms=1)


class FakeVectorSearch:
    def __init__(self) -> None:
        self.calls: list[tuple[str, int, int]] = []

    async def search(
        self, vector: list[float], kb_id: str, version: int, limit: int
    ) -> list[dict[str, object]]:
        self.calls.append((kb_id, version, limit))
        return [
            {
                "score": 0.9,
                "payload": {
                    "document_id": str(uuid.uuid5(uuid.NAMESPACE_URL, kb_id)),
                    "display_name": f"Policy {kb_id[:6]}",
                    "page_start": 4,
                    "page_end": 4,
                    "text": "The pump isolation policy requires lockout verification.",
                },
            }
        ]


def stored_text(
    session: Session, tmp_path: Path, workspace: Workspace, name: str, text: str
) -> StoredFile:
    relative = Path("workspaces") / str(workspace.id) / "uploads" / name
    path = tmp_path / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    stored = StoredFile(
        workspace_id=workspace.id,
        display_name=name,
        storage_key=relative.as_posix(),
        media_type="text/plain",
        size_bytes=path.stat().st_size,
        sha256=uuid.uuid4().hex.ljust(64, "0"),
    )
    session.add(stored)
    session.commit()
    return stored


def test_classification_is_deterministic_and_context_aware() -> None:
    assert classify_task("Fix this repository bug", False, False).agent_profile == "coding_agent"
    assert classify_task("Prepare an approval", False, True).task_type == "rag"
    assert classify_task("Inspect this file", True, False).task_type == "document_analysis"
    assert classify_task("Give a short answer", False, False).task_type == "general"


def test_model_registration_persists_capabilities_and_rejects_duplicates(
    session: Session,
) -> None:
    payload = ModelCreate(
        name="Local Qwen",
        model_key="qwen3:4b",
        capabilities=["text", "reasoning", "text"],
        context_window=32768,
        quantization="Q4_K_M",
        priority=75,
    )
    model = register_model(session, payload)

    assert session.get(RegisteredModel, model.id) is not None
    assert model.capabilities == ["text", "reasoning"]
    with pytest.raises(AppError) as duplicate:
        register_model(session, payload)
    assert duplicate.value.code == "MODEL_ALREADY_REGISTERED"


def test_disabled_model_remains_listed_but_cannot_be_selected(session: Session) -> None:
    model = register_model(
        session,
        ModelCreate(
            name="Switchable model",
            model_key="llama3.2:3b",
            capabilities=["text", "reasoning", "general"],
            priority=200,
        ),
    )

    disabled = set_model_enabled(session, model.id, False)

    assert disabled.enabled is False
    assert [item.id for item in list_models(session)] == [model.id]
    with pytest.raises(AppError) as unavailable:
        get_model(session, model.id)
    assert unavailable.value.code == "MODEL_DISABLED"

    enabled = set_model_enabled(session, model.id, True)
    assert enabled.enabled is True
    assert get_model(session, model.id).id == model.id


@pytest.mark.asyncio
async def test_multiple_workspace_files_use_bounded_temporary_rag(
    session: Session, tmp_path: Path
) -> None:
    _, workspace = tenant(session, "HybridFiles")
    pump = stored_text(
        session,
        tmp_path,
        workspace,
        "pump.txt",
        "Pump isolation requires valve V-14 to be closed. " * 40,
    )
    electrical = stored_text(
        session,
        tmp_path,
        workspace,
        "electrical.txt",
        "Electrical maintenance requires a signed permit. " * 40,
    )
    settings = Settings(
        data_root=tmp_path,
        chunk_size_chars=500,
        chunk_overlap_chars=50,
        task_ephemeral_max_chunks=16,
        task_evidence_max_hits=6,
        task_evidence_max_chars=5_000,
    )

    result = await gather_hybrid_evidence(
        session,
        workspace.id,
        [str(pump.id), str(electrical.id)],
        [],
        "What is the pump isolation procedure?",
        FakeEmbeddingProvider(),
        settings,
    )

    assert result.mode == "temporary_rag"
    assert {item["file_id"] for item in result.processed_files} == {
        str(pump.id),
        str(electrical.id),
    }
    assert result.candidate_chunks > len(result.evidence)
    assert result.evidence[0]["display_name"] == "pump.txt"
    assert all(item["retrieval_mode"] == "temporary_rag" for item in result.evidence)
    assert sum(len(str(item["text"])) for item in result.evidence) <= 5_000


@pytest.mark.asyncio
async def test_small_file_and_every_knowledge_base_are_merged(
    session: Session, tmp_path: Path
) -> None:
    _, workspace = tenant(session, "HybridKnowledge")
    uploaded = stored_text(
        session,
        tmp_path,
        workspace,
        "inspection.txt",
        "The inspection found the pump guard secured.",
    )
    bases = [
        KnowledgeBase(
            workspace_id=workspace.id,
            name=name,
            status="ready",
            active_index_version=index,
        )
        for index, name in enumerate(("Safety", "Maintenance"), start=1)
    ]
    session.add_all(bases)
    session.commit()
    vector_store = FakeVectorSearch()
    settings = Settings(data_root=tmp_path, task_direct_read_chars=10_000)

    result = await gather_hybrid_evidence(
        session,
        workspace.id,
        [str(uploaded.id)],
        [str(base.id) for base in bases],
        "What is the pump isolation procedure?",
        FakeEmbeddingProvider(),
        settings,
        vector_store,
    )

    assert result.mode == "hybrid"
    assert result.processed_files[0]["mode"] == "direct_read"
    assert {call[0] for call in vector_store.calls} == {str(base.id) for base in bases}
    assert {item["source_type"] for item in result.evidence} == {
        "workspace_file",
        "knowledge_base",
    }
    assert [item["source_id"] for item in result.evidence] == [
        f"S{index}" for index in range(1, len(result.evidence) + 1)
    ]


def test_task_creation_is_durable_idempotent_and_tenant_scoped(session: Session) -> None:
    owner, workspace = tenant(session, "Alpha")
    outsider, _ = tenant(session, "Beta")
    payload = TaskCreate(
        workspace_id=workspace.id,
        goal="Prepare a local approval recommendation",
        requested_outputs=["docx"],
    )
    settings = Settings(task_timeout_seconds=60)

    first = create_task(session, payload, owner, settings, "stable-key")
    repeated = create_task(session, payload, owner, settings, "stable-key")

    assert repeated.task_id == first.task_id
    assert session.scalar(select(TaskRun).where(TaskRun.id == first.run_id)) is not None
    assert len(list(session.scalars(select(Task)))) == 1
    with pytest.raises(AppError) as denied:
        get_owned_task(session, first.task_id, outsider)
    assert denied.value.code == "TASK_NOT_FOUND"


def test_tool_policy_and_validated_docx(tmp_path: Path) -> None:
    registry = ToolRegistry()
    registry.authorize("general_agent", "create_docx")
    with pytest.raises(AppError):
        registry.authorize("coding_agent", "search_knowledge")

    published = create_approval_docx(
        tmp_path,
        uuid.uuid4(),
        uuid.uuid4(),
        "Approve the pump procedure",
        "## Decision\n\n**Approve** after a human verifies isolation. [S1]\n\n"
        "- Confirm lockout\n- Record sign-off\n\n"
        "| Control | Status |\n| --- | --- |\n| Isolation | Ready |",
        [{"source_id": "S1", "display_name": "Pump SOP", "page_start": 4, "page_end": 4}],
        "Acme Engineering",
    )
    artifact = tmp_path / published.storage_key
    document = DocxDocument(str(artifact))

    assert artifact.is_file()
    assert published.size_bytes == artifact.stat().st_size
    assert "Approval Control" in {paragraph.text for paragraph in document.paragraphs}
    assert "Acme Engineering™" in document.sections[0].header.paragraphs[0].text
    assert "Acme Engineering™" in document.sections[0].footer.paragraphs[0].text
    assert "**" not in "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert any(
        run.bold and run.text == "Approve"
        for paragraph in document.paragraphs
        for run in paragraph.runs
    )
    assert any(paragraph.style.name == "List Bullet" for paragraph in document.paragraphs)
    assert document.tables[0].cell(1, 0).text == "Isolation"


@pytest.mark.asyncio
async def test_read_file_extracts_a_new_pdf_without_vector_indexing(
    session: Session, tmp_path: Path
) -> None:
    _, workspace = tenant(session, "PDF")
    relative = Path("workspaces") / str(workspace.id) / "uploads" / "manual.pdf"
    source = tmp_path / relative
    source.parent.mkdir(parents=True)
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Valve V-14 must be closed before pump maintenance.")
    pdf.save(source)
    pdf.close()
    stored = StoredFile(
        workspace_id=workspace.id,
        display_name="manual.pdf",
        storage_key=relative.as_posix(),
        media_type="application/pdf",
        size_bytes=source.stat().st_size,
        sha256="0" * 64,
    )
    session.add(stored)
    session.commit()
    settings = Settings(data_root=tmp_path)

    result = await ToolRegistry().execute(
        "document_agent",
        "read_file",
        {"file_id": str(stored.id)},
        ToolContext(
            workspace_id=workspace.id,
            session=session,
            settings=settings,
            provider=OllamaModelProvider("http://unused"),
        ),
    )

    extracted = session.scalar(select(Document).where(Document.file_id == stored.id))
    assert "Valve V-14" in result.data["text"]
    assert result.data["page_count"] == 1
    assert extracted is not None
    assert extracted.extraction_status == "completed"
    assert extracted.normalized_path is not None
